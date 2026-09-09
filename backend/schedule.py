"""
Ironmongery scheduling: products, hardware sets, door types, doors, and the
schedule output. Kept out of main.py so the counting app and the scheduling
app can grow without tangling.
"""
import io
import re
import uuid
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import FileResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

import models, auth
from database import get_db

router = APIRouter(prefix="/api", tags=["schedule"])

PRODUCT_IMG_DIR = Path("uploads") / "products"
PRODUCT_IMG_DIR.mkdir(parents=True, exist_ok=True)


# ── Schemas ───────────────────────────────────────────────────────────────────
class ProductOut(BaseModel):
    id: int; sku: str; name: str; category: str; unit: str
    cost: Optional[float] = None; sell: Optional[float] = None
    intec_code: str = ""; image_url: str = ""; notes: str = ""; active: bool = True
    used_in: list[str] = []

class ProductIn(BaseModel):
    sku: str; name: str; category: str = "Other"; unit: str = "EACH"
    cost: Optional[float] = None; sell: Optional[float] = None
    intec_code: str = ""; notes: str = ""; active: bool = True

class SetItemIn(BaseModel):
    product_id: int; qty: int = 1

class SetIn(BaseModel):
    code: str; name: str; description: str = ""; fire_rated: bool = False
    notes: str = ""; items: list[SetItemIn] = []

class SetItemOut(BaseModel):
    id: int; product_id: int; sku: str; name: str; category: str; unit: str
    qty: int; cost: Optional[float] = None; image_url: str = ""

class SetUse(BaseModel):
    project_id: int; project: str; doors: int

class SetOut(BaseModel):
    id: int; code: str; name: str; description: str = ""; fire_rated: bool = False
    notes: str = ""; archived: bool = False; copied_from: Optional[str] = None
    product_count: int = 0; items_per_door: int = 0; cost_per_door: Optional[float] = None
    used_on: list[SetUse] = []; items: list[SetItemOut] = []


# ── Helpers ───────────────────────────────────────────────────────────────────
def _img_url(p: models.Product) -> str:
    return f"/api/files/products/{Path(p.image_path).name}" if p.image_path else ""


def _used_map(db: Session) -> dict[int, list[str]]:
    """product_id -> [set codes] for every live set."""
    out: dict[int, list[str]] = {}
    rows = (db.query(models.SetItem.product_id, models.HardwareSet.code)
              .join(models.HardwareSet, models.HardwareSet.id == models.SetItem.set_id)
              .filter(models.HardwareSet.archived == False).all())   # noqa: E712
    for pid, code in rows:
        out.setdefault(pid, [])
        if code not in out[pid]:
            out[pid].append(code)
    return out


def _product_out(p: models.Product, used: dict) -> ProductOut:
    return ProductOut(
        id=p.id, sku=p.sku, name=p.name, category=p.category or "Other", unit=p.unit or "EACH",
        cost=p.cost, sell=p.sell, intec_code=p.intec_code or "", image_url=_img_url(p),
        notes=p.notes or "", active=bool(p.active), used_in=sorted(used.get(p.id, [])),
    )


def _set_uses(db: Session, sid: int) -> list[SetUse]:
    """Projects using this set, with door counts (type assignment or per-door override)."""
    counts: dict[int, int] = {}
    names: dict[int, str] = {}
    for d in db.query(models.Door).all():
        eff = d.set_id or (d.door_type.set_id if d.door_type else None)
        if eff == sid:
            counts[d.project_id] = counts.get(d.project_id, 0) + 1
    if counts:
        for p in db.query(models.Project).filter(models.Project.id.in_(list(counts))).all():
            names[p.id] = p.name
    return [SetUse(project_id=pid, project=names.get(pid, f"Project {pid}"), doors=n)
            for pid, n in sorted(counts.items())]


def _set_out(db: Session, s: models.HardwareSet, with_uses: bool = True) -> SetOut:
    items = []
    cost_total, cost_known = 0.0, True
    for it in s.items:
        p = it.product
        items.append(SetItemOut(id=it.id, product_id=p.id, sku=p.sku, name=p.name,
                                category=p.category or "Other", unit=p.unit or "EACH",
                                qty=it.qty, cost=p.cost, image_url=_img_url(p)))
        if p.cost is None:
            cost_known = False
        else:
            cost_total += p.cost * it.qty
    parent = db.query(models.HardwareSet).get(s.copied_from_id) if s.copied_from_id else None
    return SetOut(
        id=s.id, code=s.code, name=s.name, description=s.description or "",
        fire_rated=bool(s.fire_rated), notes=s.notes or "", archived=bool(s.archived),
        copied_from=f"{parent.code} {parent.name}" if parent else None,
        product_count=len(items), items_per_door=sum(i.qty for i in items),
        cost_per_door=round(cost_total, 2) if (items and cost_known) else None,
        used_on=_set_uses(db, s.id) if with_uses else [], items=items,
    )


def _next_set_code(db: Session) -> str:
    best = 0
    for (code,) in db.query(models.HardwareSet.code).all():
        m = re.fullmatch(r"MF\s*(\d+)", (code or "").strip(), re.I)
        if m:
            best = max(best, int(m.group(1)))
    return f"MF {best + 1:02d}"


# ── Products ──────────────────────────────────────────────────────────────────
@router.get("/products", response_model=list[ProductOut])
def list_products(q: str = "", category: str = "", include_inactive: bool = False,
                  db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    qry = db.query(models.Product)
    if not include_inactive:
        qry = qry.filter(models.Product.active == True)   # noqa: E712
    if category:
        qry = qry.filter(models.Product.category == category)
    if q:
        like = f"%{q.strip()}%"
        qry = qry.filter((models.Product.sku.ilike(like)) | (models.Product.name.ilike(like))
                         | (models.Product.intec_code.ilike(like)))
    used = _used_map(db)
    return [_product_out(p, used) for p in qry.order_by(models.Product.sku).all()]


@router.get("/products/categories")
def product_categories(db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    rows = (db.query(models.Product.category, models.Product.id)
              .filter(models.Product.active == True).all())   # noqa: E712
    counts: dict[str, int] = {}
    for cat, _ in rows:
        counts[cat or "Other"] = counts.get(cat or "Other", 0) + 1
    return [{"name": k, "count": v} for k, v in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))]


@router.post("/products/import")
async def import_products(file: UploadFile = File(...), db: Session = Depends(get_db),
                          cu=Depends(auth.get_current_user)):
    """
    Load the Cin7 'Products Price List' export (.xlsx). Finds the header row by
    the SKU column, upserts by SKU. Existing photos, Intec codes and notes are kept.
    """
    import openpyxl
    raw = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(400, "That file is not an Excel workbook (.xlsx)")
    ws = wb.worksheets[0]
    rows = [r for r in ws.iter_rows(values_only=True)]
    hi = next((i for i, r in enumerate(rows)
               if r and any(str(c).strip().lower() == "sku" for c in r if c is not None)), None)
    if hi is None:
        raise HTTPException(400, "Could not find a SKU column. Export the Products Price List report from Cin7.")
    hdr = [str(c).strip().lower() if c is not None else "" for c in rows[hi]]
    def col(*names):
        for n in names:
            if n in hdr:
                return hdr.index(n)
        return None
    c_sku, c_name, c_unit, c_cat = col("sku"), col("product", "name"), col("unit"), col("category")
    c_cost, c_gen, c_whole = col("average cost", "averagecost", "cost"), col("general"), col("wholesale")
    if c_sku is None or c_name is None:
        raise HTTPException(400, "The export needs at least SKU and Product columns.")

    def num(v):
        try:
            f = float(v)
            return f if f > 0 else None
        except (TypeError, ValueError):
            return None

    added = updated = skipped = 0
    for r in rows[hi + 1:]:
        if not r or c_sku >= len(r) or r[c_sku] in (None, ""):
            continue
        sku = str(r[c_sku]).strip()
        name = str(r[c_name]).strip() if c_name < len(r) and r[c_name] is not None else ""
        if not sku or not name or sku.lower() in ("total", "grand total"):
            skipped += 1
            continue
        cat = (str(r[c_cat]).strip() if c_cat is not None and r[c_cat] else "Other") or "Other"
        unit = (str(r[c_unit]).strip() if c_unit is not None and r[c_unit] else "EACH") or "EACH"
        cost = num(r[c_cost]) if c_cost is not None else None
        if not cost:
            cost = None            # Cin7 shows 0 until the first delivery lands
        sell = None
        for ci in (c_whole, c_gen):
            if ci is not None and num(r[ci]):
                sell = num(r[ci]); break
        p = db.query(models.Product).filter(models.Product.sku == sku).first()
        if p:
            p.name, p.category, p.unit = name, cat, unit.upper()
            if cost is not None: p.cost = cost
            if sell is not None: p.sell = sell
            updated += 1
        else:
            db.add(models.Product(sku=sku, name=name, category=cat, unit=unit.upper(),
                                  cost=cost, sell=sell, source="cin7"))
            added += 1
    db.commit()
    return {"added": added, "updated": updated, "skipped": skipped}


@router.post("/products", response_model=ProductOut, status_code=201)
def create_product(payload: ProductIn, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    sku = payload.sku.strip()
    if not sku or not payload.name.strip():
        raise HTTPException(400, "Code and name are required")
    if db.query(models.Product).filter(models.Product.sku == sku).first():
        raise HTTPException(409, f"A product with code {sku} already exists")
    p = models.Product(**{**payload.model_dump(), "sku": sku, "name": payload.name.strip(), "source": "manual"})
    db.add(p); db.commit(); db.refresh(p)
    return _product_out(p, {})


@router.put("/products/{pid}", response_model=ProductOut)
def update_product(pid: int, payload: ProductIn, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    p = db.query(models.Product).get(pid)
    if not p:
        raise HTTPException(404, "Product not found")
    sku = payload.sku.strip()
    clash = db.query(models.Product).filter(models.Product.sku == sku, models.Product.id != pid).first()
    if clash:
        raise HTTPException(409, f"Another product already uses code {sku}")
    for k, v in payload.model_dump().items():
        setattr(p, k, v.strip() if isinstance(v, str) else v)
    db.commit(); db.refresh(p)
    return _product_out(p, _used_map(db))


@router.post("/products/{pid}/image", response_model=ProductOut)
async def upload_product_image(pid: int, file: UploadFile = File(...),
                               db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    p = db.query(models.Product).get(pid)
    if not p:
        raise HTTPException(404, "Product not found")
    ext = Path(file.filename or "").suffix.lower()
    if ext not in (".png", ".jpg", ".jpeg", ".webp", ".gif"):
        raise HTTPException(400, "Use a PNG, JPG, WEBP or GIF image")
    dest = PRODUCT_IMG_DIR / f"{uuid.uuid4().hex}{ext}"
    dest.write_bytes(await file.read())
    if p.image_path and Path(p.image_path).exists():
        try: Path(p.image_path).unlink()
        except OSError: pass
    p.image_path = str(dest)
    db.commit(); db.refresh(p)
    return _product_out(p, _used_map(db))


@router.delete("/products/{pid}", status_code=204)
def archive_product(pid: int, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    p = db.query(models.Product).get(pid)
    if not p:
        raise HTTPException(404, "Product not found")
    p.active = False
    db.commit()


@router.get("/files/products/{filename}")
def serve_product_image(filename: str, cu=Depends(auth.get_current_user)):
    path = PRODUCT_IMG_DIR / Path(filename).name
    if not path.exists():
        raise HTTPException(404, "Image not found")
    return FileResponse(str(path))


# ── Hardware sets ─────────────────────────────────────────────────────────────
@router.get("/sets", response_model=list[SetOut])
def list_sets(include_archived: bool = False, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    qry = db.query(models.HardwareSet)
    if not include_archived:
        qry = qry.filter(models.HardwareSet.archived == False)   # noqa: E712
    return [_set_out(db, s) for s in qry.order_by(models.HardwareSet.code).all()]


@router.get("/sets/next-code")
def next_set_code(db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    return {"code": _next_set_code(db)}


@router.get("/sets/{sid}", response_model=SetOut)
def get_set(sid: int, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    s = db.query(models.HardwareSet).get(sid)
    if not s:
        raise HTTPException(404, "Set not found")
    return _set_out(db, s)


def _apply_items(db: Session, s: models.HardwareSet, items: list[SetItemIn]):
    seen = set()
    s.items.clear()
    db.flush()
    for i, it in enumerate(items):
        if it.product_id in seen or it.qty <= 0:
            continue
        if not db.query(models.Product).get(it.product_id):
            raise HTTPException(400, f"Product {it.product_id} does not exist")
        seen.add(it.product_id)
        s.items.append(models.SetItem(product_id=it.product_id, qty=it.qty, sort_order=i))


@router.post("/sets", response_model=SetOut, status_code=201)
def create_set(payload: SetIn, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    code = payload.code.strip() or _next_set_code(db)
    if not payload.name.strip():
        raise HTTPException(400, "A set needs a name")
    if db.query(models.HardwareSet).filter(models.HardwareSet.code == code,
                                           models.HardwareSet.archived == False).first():   # noqa: E712
        raise HTTPException(409, f"Set {code} already exists")
    s = models.HardwareSet(code=code, name=payload.name.strip(), description=payload.description,
                           fire_rated=payload.fire_rated, notes=payload.notes, created_by_id=cu.id)
    db.add(s); db.flush()
    _apply_items(db, s, payload.items)
    db.commit(); db.refresh(s)
    return _set_out(db, s)


@router.put("/sets/{sid}", response_model=SetOut)
def update_set(sid: int, payload: SetIn, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    s = db.query(models.HardwareSet).get(sid)
    if not s:
        raise HTTPException(404, "Set not found")
    code = payload.code.strip() or s.code
    clash = db.query(models.HardwareSet).filter(models.HardwareSet.code == code, models.HardwareSet.id != sid,
                                                models.HardwareSet.archived == False).first()   # noqa: E712
    if clash:
        raise HTTPException(409, f"Another set already uses code {code}")
    s.code, s.name = code, payload.name.strip() or s.name
    s.description, s.fire_rated, s.notes = payload.description, payload.fire_rated, payload.notes
    _apply_items(db, s, payload.items)
    db.commit(); db.refresh(s)
    return _set_out(db, s)


@router.post("/sets/{sid}/copy", response_model=SetOut, status_code=201)
def copy_set(sid: int, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    src = db.query(models.HardwareSet).get(sid)
    if not src:
        raise HTTPException(404, "Set not found")
    s = models.HardwareSet(code=_next_set_code(db), name=f"{src.name} (copy)", description=src.description,
                           fire_rated=src.fire_rated, notes=src.notes, copied_from_id=src.id, created_by_id=cu.id)
    db.add(s); db.flush()
    for i, it in enumerate(src.items):
        s.items.append(models.SetItem(product_id=it.product_id, qty=it.qty, sort_order=i))
    db.commit(); db.refresh(s)
    return _set_out(db, s)


@router.delete("/sets/{sid}", status_code=204)
def archive_set(sid: int, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    s = db.query(models.HardwareSet).get(sid)
    if not s:
        raise HTTPException(404, "Set not found")
    if _set_uses(db, sid):
        raise HTTPException(409, "This set is assigned to doors on a project. Reassign them first.")
    s.archived = True
    db.commit()


# ═════════════════════════════════════════════════════════════════════════════
# Doors
# ═════════════════════════════════════════════════════════════════════════════
UPLOAD_DIR = Path("uploads")

DOOR_STATUSES = ("decide", "assigned", "excluded")


class FloorCount(BaseModel):
    floor: str; count: int

class Suggestion(BaseModel):
    set_id: int; set_code: str; set_name: str; reason: str

class DoorTypeOut(BaseModel):
    id: int; code: str; description: str = ""; fire_rating: str = ""; acoustic: str = ""
    width: Optional[int] = None; height: Optional[int] = None; spec_text: str = ""
    status: str = "decide"; set_id: Optional[int] = None; set_code: str = ""; set_name: str = ""
    sort_order: int = 0; door_count: int = 0; handed_count: int = 0
    floors: list[FloorCount] = []; suggestion: Optional[Suggestion] = None

class DoorTypeIn(BaseModel):
    code: str; description: str = ""; fire_rating: str = ""; acoustic: str = ""
    width: Optional[int] = None; height: Optional[int] = None; spec_text: str = ""
    status: str = "decide"; set_id: Optional[int] = None

class DoorOut(BaseModel):
    id: int; ref: str; floor: str; handed: bool = False
    door_type_id: Optional[int] = None; type_code: str = ""
    drawing_id: Optional[int] = None; drawing: str = ""; page_number: int = 1
    x: Optional[float] = None; y: Optional[float] = None
    set_id: Optional[int] = None; effective_set_id: Optional[int] = None; effective_set_code: str = ""
    source: str = "plan"; note: str = ""

class DoorIn(BaseModel):
    ref: str = ""; floor: str = ""; handed: bool = False
    door_type_id: Optional[int] = None; set_id: Optional[int] = None; note: str = ""

class PlanOut(BaseModel):
    id: int; name: str; floor: str = ""; doors: int = 0; status: str = ""; error: str = ""

class BulkDoorsIn(BaseModel):
    door_ids: list[int] = []; untyped: bool = False
    set_id: Optional[int] = None; door_type_id: Optional[int] = None; clear_set: bool = False

class DoorsSummary(BaseModel):
    project_id: int; kind: str = "symbols"; doors: int; plans: int; untyped_doors: int
    untyped_floors: list[FloorCount] = []; untyped_with_set: int = 0; untyped_sample: list[str] = []
    decide: int; assigned: int; excluded: int
    assigned_doors: int; suggestions: int; sets_available: int = 0
    floors: list[str] = []; door_types: list[DoorTypeOut] = []; plan_list: list[PlanOut] = []


# ── Helpers ───────────────────────────────────────────────────────────────────
def _own_project(pid: int, db: Session, cu) -> models.Project:
    p = db.query(models.Project).filter(models.Project.id == pid, models.Project.owner_id == cu.id).first()
    if not p:
        raise HTTPException(404, "Project not found")
    return p


def _own_door_type(dtid: int, db: Session, cu) -> models.DoorType:
    dt = db.query(models.DoorType).get(dtid)
    if not dt:
        raise HTTPException(404, "Door type not found")
    _own_project(dt.project_id, db, cu)
    return dt


def _norm_code(code: str) -> str:
    c = re.sub(r"\s+", "", (code or "").upper())
    m = re.fullmatch(r"([A-Z]{1,3})-?(\d{1,3})(?:[-.](\d{1,2}))?([A-Z]?)", c)
    if m:
        num = m.group(2)
        num = num.zfill(2) if len(num) < 3 else num
        return f"{m.group(1)}-{num}" + (f"-{int(m.group(3)):02d}" if m.group(3) else "") + m.group(4)
    return c


def _code_sort(code: str) -> tuple:
    m = re.search(r"(\d+)", code or "")
    return (code.split(str(m.group(1)))[0] if m else code, int(m.group(1)) if m else 0, code)


_STOP = {"door", "doors", "internal", "int", "the", "a", "and", "to", "of", "with", "sgl", "single",
         "dbl", "double", "fr", "nfr", "fire", "rated", "leaf"}


def _tokens(s: str) -> set:
    return {t for t in re.findall(r"[a-z0-9]+", (s or "").lower()) if t not in _STOP and len(t) > 1}


def _suggest_for(db: Session, dt: models.DoorType, sets: list, foreign_types: list) -> Optional[Suggestion]:
    """
    A set for a door type that still needs deciding, from what was done before:
    the same description on another job, the same code with a matching
    description, or a set whose name says the same thing.
    """
    desc_tokens = _tokens(dt.description)
    code = _norm_code(dt.code)
    fire = (dt.fire_rating or "").upper().startswith("FD")
    # 1. Same description used on another job
    best = None
    for ft, proj_name in foreign_types:
        if not ft.set_id or ft.hardware_set is None or ft.hardware_set.archived:
            continue
        same_desc = desc_tokens and _tokens(ft.description) == desc_tokens
        same_code = _norm_code(ft.code) == code
        if same_desc:
            score = 3 if same_code else 2
        elif same_code and desc_tokens and _tokens(ft.description) & desc_tokens:
            score = 1
        else:
            continue
        if best is None or score > best[0]:
            best = (score, ft, proj_name)
    if best:
        _, ft, proj_name = best
        s = ft.hardware_set
        return Suggestion(set_id=s.id, set_code=s.code, set_name=s.name,
                          reason=f"Used for {ft.code} {ft.description} on {proj_name}".strip())
    # 2. A set whose name matches the description
    if desc_tokens:
        scored = []
        for s in sets:
            st = _tokens(s.name)
            hit = len(st & desc_tokens)
            if hit == 0:
                continue
            frac = hit / len(desc_tokens)
            if frac < 0.5:
                continue
            scored.append((frac + (0.1 if bool(s.fire_rated) == fire else 0), s))
        if scored:
            scored.sort(key=lambda t: -t[0])
            s = scored[0][1]
            return Suggestion(set_id=s.id, set_code=s.code, set_name=s.name, reason="Set name matches this door type")
    return None


def _door_type_out(db: Session, dt: models.DoorType, doors: list, sets=None, foreign_types=None) -> DoorTypeOut:
    floors: dict[str, int] = {}
    for d in doors:
        floors[d.floor or "Unknown"] = floors.get(d.floor or "Unknown", 0) + 1
    s = dt.hardware_set if dt.set_id else None
    if s and s.archived:
        s = None
    sug = None
    if dt.status == "decide" and not dt.set_id and sets is not None:
        sug = _suggest_for(db, dt, sets, foreign_types or [])
    return DoorTypeOut(
        id=dt.id, code=dt.code, description=dt.description or "", fire_rating=dt.fire_rating or "",
        acoustic=dt.acoustic or "", width=dt.width, height=dt.height, spec_text=dt.spec_text or "",
        status=dt.status or "decide", set_id=s.id if s else None,
        set_code=s.code if s else "", set_name=s.name if s else "", sort_order=dt.sort_order or 0,
        door_count=len(doors), handed_count=sum(1 for d in doors if d.handed),
        floors=[FloorCount(floor=f, count=n) for f, n in sorted(floors.items(), key=lambda t: _floor_rank(t[0]))],
        suggestion=sug,
    )


_FLOOR_ORDER = ["Basement", "Lower Ground", "Ground", "Mezzanine", "First", "Second", "Third", "Fourth",
                "Fifth", "Sixth", "Seventh", "Eighth", "Ninth", "Tenth", "Penthouse", "Roof"]


def _floor_rank(f: str):
    if f in _FLOOR_ORDER:
        return (0, _FLOOR_ORDER.index(f), f)
    m = re.match(r"Level (\d+)", f or "")
    if m:
        return (1, int(m.group(1)), f)
    return (2, 0, f or "")


def _door_out(d: models.Door, dwg_names: dict, set_codes: dict) -> DoorOut:
    eff = d.set_id or (d.door_type.set_id if d.door_type else None)
    return DoorOut(
        id=d.id, ref=d.ref or "", floor=d.floor or "", handed=bool(d.handed),
        door_type_id=d.door_type_id, type_code=d.door_type.code if d.door_type else "",
        drawing_id=d.drawing_id, drawing=dwg_names.get(d.drawing_id, ""), page_number=d.page_number or 1,
        x=d.x, y=d.y, set_id=d.set_id, effective_set_id=eff, effective_set_code=set_codes.get(eff, ""),
        source=d.source or "plan", note=d.note or "",
    )


def _get_or_create_type(db: Session, pid: int, code: str, cache: dict) -> models.DoorType:
    key = _norm_code(code)
    if key in cache:
        return cache[key]
    dt = models.DoorType(project_id=pid, code=key, status="decide", sort_order=_code_sort(key)[1])
    # An earlier schedule import may have left DT-05-01 / DT-05-02 as their own
    # types with no doors; the plans tag DT-05, so fold them into this one.
    for vkey in sorted(k for k in list(cache) if k.startswith(key + "-")):
        v = cache[vkey]
        if db.query(models.Door).filter(models.Door.door_type_id == v.id).count():
            continue
        vd = v.description or ""
        desc = re.sub(r"\s*\(.*?\)?\s*$", "", vd).strip() if "(" in vd else vd
        if desc and not dt.description: dt.description = desc
        if v.fire_rating and not dt.fire_rating: dt.fire_rating = v.fire_rating
        if v.acoustic and not dt.acoustic: dt.acoustic = v.acoustic
        if v.width and not dt.width: dt.width = v.width
        if v.height and not dt.height: dt.height = v.height
        if not dt.set_id and v.set_id: dt.set_id = v.set_id; dt.status = v.status
        line = f"{vkey}: {vd}" + (f" {v.width} x {v.height}" if v.width and v.height else "")
        dt.spec_text = ((dt.spec_text or "").rstrip() + "\n" + line).strip()
        db.delete(v); del cache[vkey]
    db.add(dt); db.flush()
    cache[key] = dt
    return dt


def _type_cache(db: Session, pid: int) -> dict:
    return {_norm_code(t.code): t for t in db.query(models.DoorType).filter(models.DoorType.project_id == pid).all()}


def sync_doors_from_drawing(db: Session, drawing: models.Drawing, pdf_path: Optional[str] = None) -> int:
    """
    Read the door tags off every page of a drawing and (re)create its doors.
    Plan-sourced doors for the drawing are replaced; typed or imported doors
    are left alone. Returns the number of doors found.
    """
    from detection.doors import extract_door_tags, floor_from_name, floor_short, looks_like_door_plan, split_types_and_refs
    path = Path(pdf_path) if pdf_path else UPLOAD_DIR / drawing.filename / "drawing.pdf"
    if not path.exists():
        return 0
    db.query(models.Door).filter(models.Door.drawing_id == drawing.id,
                                 models.Door.source == "plan").delete(synchronize_session=False)
    floor = floor_from_name(drawing.original_name) or floor_from_name(drawing.level or "") or (drawing.level or "")
    cache = _type_cache(db, drawing.project_id)
    # Read every page first: whether "D-01" is a type or a door number depends on repetition
    pages: list[tuple[int, list]] = []
    for page_no in range(1, (drawing.total_pages or 1) + 1):
        try:
            tags = extract_door_tags(str(path), page_no)
        except Exception:
            continue
        if looks_like_door_plan(tags):
            pages.append((page_no, tags))
    all_tags = [t for _, tags in pages for t in tags]
    split_types_and_refs(all_tags, set(cache))
    if not floor:
        hints = [t["floor_hint"] for t in all_tags if t.get("floor_hint")]
        if hints:
            floor = max(set(hints), key=hints.count)
    counters: dict[tuple, int] = {}
    # Existing generated refs on other drawings keep numbering unique per floor and type
    for (ref,) in db.query(models.Door.ref).filter(models.Door.project_id == drawing.project_id).all():
        m = re.fullmatch(r"([A-Z0-9]{1,3})-([A-Z0-9-]+)-(\d+)", ref or "")
        if m:
            key = (m.group(1), m.group(2))
            counters[key] = max(counters.get(key, 0), int(m.group(3)))
    total = 0
    for page_no, tags in pages:
        tags.sort(key=lambda t: (round(t["y"], 2), t["x"]))
        for t in tags:
            code = t["type_code"]
            ref = t["ref"]
            # A sheet can carry two floors (ground and first of a house); the label
            # under the tag says which, and beats the sheet-wide floor.
            door_floor = t.get("floor_hint") or floor
            fs = floor_short(door_floor)
            dt = _get_or_create_type(db, drawing.project_id, code, cache) if code else None
            if not ref:
                short = code.replace("DT-", "") if code else "D"
                counters[(fs, short)] = counters.get((fs, short), 0) + 1
                ref = f"{fs}-{short}-{counters[(fs, short)]:02d}"
            db.add(models.Door(project_id=drawing.project_id, door_type_id=dt.id if dt else None,
                               ref=ref, floor=door_floor, handed=t["handed"], drawing_id=drawing.id,
                               page_number=page_no, x=t["x"], y=t["y"], source="plan"))
            total += 1
    db.commit()
    return total


# ── Endpoints ─────────────────────────────────────────────────────────────────
@router.get("/projects/{pid}/doors/summary", response_model=DoorsSummary)
def doors_summary(pid: int, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    project = _own_project(pid, db, cu)
    types = db.query(models.DoorType).filter(models.DoorType.project_id == pid).all()
    doors = db.query(models.Door).filter(models.Door.project_id == pid).all()
    by_dwg: dict[int, int] = {}
    for d in doors:
        if d.drawing_id:
            by_dwg[d.drawing_id] = by_dwg.get(d.drawing_id, 0) + 1
    from detection.doors import floor_from_name
    plan_list = [PlanOut(id=dw.id, name=dw.level or dw.original_name,
                         floor=floor_from_name(dw.original_name) or floor_from_name(dw.level or ""),
                         doors=by_dwg.get(dw.id, 0), status=dw.status or "", error=dw.error_message or "")
                 for dw in sorted(project.drawings, key=lambda x: x.id)]
    by_type: dict[int, list] = {}
    for d in doors:
        by_type.setdefault(d.door_type_id, []).append(d)
    sets = db.query(models.HardwareSet).filter(models.HardwareSet.archived == False).all()   # noqa: E712
    foreign = [(ft, ft.project.name if ft.project else "another job")
               for ft in db.query(models.DoorType).filter(models.DoorType.project_id != pid,
                                                          models.DoorType.set_id.isnot(None)).all()]
    outs = [_door_type_out(db, t, by_type.get(t.id, []), sets, foreign) for t in types]
    outs.sort(key=lambda o: _code_sort(o.code))
    floors = sorted({d.floor for d in doors if d.floor}, key=_floor_rank)
    assigned_doors = sum(1 for d in doors if d.set_id or (d.door_type and d.door_type.set_id
                                                          and d.door_type.status != "excluded"))
    untyped = by_type.get(None, [])
    uf: dict[str, int] = {}
    for d in untyped:
        uf[d.floor or "Unknown"] = uf.get(d.floor or "Unknown", 0) + 1
    return DoorsSummary(
        untyped_floors=[FloorCount(floor=f, count=n) for f, n in sorted(uf.items(), key=lambda t: _floor_rank(t[0]))],
        untyped_with_set=sum(1 for d in untyped if d.set_id),
        untyped_sample=[d.ref for d in sorted(untyped, key=lambda d: d.ref or "")[:6]],
        project_id=pid, kind=project.kind or "symbols", sets_available=len(sets), plan_list=plan_list,
        doors=len(doors), plans=len({d.drawing_id for d in doors if d.drawing_id}),
        untyped_doors=len(by_type.get(None, [])),
        decide=sum(1 for o in outs if o.status == "decide"),
        assigned=sum(1 for o in outs if o.status == "assigned"),
        excluded=sum(1 for o in outs if o.status == "excluded"),
        assigned_doors=assigned_doors, suggestions=sum(1 for o in outs if o.suggestion),
        floors=floors, door_types=outs,
    )


@router.get("/projects/{pid}/doors/count")
def doors_count(pid: int, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    _own_project(pid, db, cu)
    n = db.query(models.Door).filter(models.Door.project_id == pid).count()
    types = db.query(models.DoorType).filter(models.DoorType.project_id == pid).all()
    return {"doors": n, "types": len(types), "decide": sum(1 for t in types if t.status == "decide")}


@router.get("/projects/{pid}/doors", response_model=list[DoorOut])
def list_doors(pid: int, type_id: Optional[int] = None, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    _own_project(pid, db, cu)
    q = db.query(models.Door).filter(models.Door.project_id == pid)
    if type_id is not None:
        q = q.filter(models.Door.door_type_id == (type_id or None))
    doors = q.all()
    dwg_names = {d.id: (d.level or d.original_name) for d in db.query(models.Drawing).filter(models.Drawing.project_id == pid).all()}
    set_codes = {s.id: s.code for s in db.query(models.HardwareSet).all()}
    doors.sort(key=lambda d: (_floor_rank(d.floor or ""), d.ref or ""))
    return [_door_out(d, dwg_names, set_codes) for d in doors]


def _apply_type_payload(dt: models.DoorType, payload: DoorTypeIn, db: Session):
    if payload.status not in DOOR_STATUSES:
        raise HTTPException(400, "Bad status")
    if payload.set_id is not None and not db.query(models.HardwareSet).get(payload.set_id):
        raise HTTPException(404, "Set not found")
    dt.code = _norm_code(payload.code) or dt.code
    dt.description = payload.description.strip(); dt.fire_rating = payload.fire_rating.strip()
    dt.acoustic = payload.acoustic.strip(); dt.width = payload.width; dt.height = payload.height
    dt.spec_text = payload.spec_text; dt.set_id = payload.set_id
    if payload.status == "excluded":
        dt.status = "excluded"
    else:
        dt.status = "assigned" if payload.set_id else "decide"
    dt.sort_order = _code_sort(dt.code)[1]


@router.post("/projects/{pid}/door-types", response_model=DoorTypeOut, status_code=201)
def create_door_type(pid: int, payload: DoorTypeIn, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    _own_project(pid, db, cu)
    code = _norm_code(payload.code)
    if not code:
        raise HTTPException(400, "Door type needs a code")
    if code in _type_cache(db, pid):
        raise HTTPException(409, f"{code} already exists on this project")
    dt = models.DoorType(project_id=pid)
    _apply_type_payload(dt, payload, db)
    db.add(dt); db.commit(); db.refresh(dt)
    return _door_type_out(db, dt, [])


@router.put("/door-types/{dtid}", response_model=DoorTypeOut)
def update_door_type(dtid: int, payload: DoorTypeIn, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    dt = _own_door_type(dtid, db, cu)
    new_code = _norm_code(payload.code)
    if new_code and new_code != _norm_code(dt.code) and new_code in _type_cache(db, dt.project_id):
        raise HTTPException(409, f"{new_code} already exists on this project")
    _apply_type_payload(dt, payload, db)
    db.commit(); db.refresh(dt)
    doors = db.query(models.Door).filter(models.Door.door_type_id == dt.id).all()
    return _door_type_out(db, dt, doors)


@router.delete("/door-types/{dtid}", status_code=204)
def delete_door_type(dtid: int, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    dt = _own_door_type(dtid, db, cu)
    n = db.query(models.Door).filter(models.Door.door_type_id == dt.id).count()
    if n:
        raise HTTPException(409, f"{dt.code} has {n} door{'s' if n != 1 else ''}. Move or remove them first.")
    db.delete(dt); db.commit()


@router.post("/projects/{pid}/doors/apply-suggestions")
def apply_suggestions(pid: int, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    """Assign every suggested set to its door type in one go."""
    summary = doors_summary(pid, db, cu)
    applied = 0
    for o in summary.door_types:
        if o.suggestion and o.status == "decide":
            dt = db.query(models.DoorType).get(o.id)
            dt.set_id = o.suggestion.set_id; dt.status = "assigned"; applied += 1
    db.commit()
    return {"applied": applied}


@router.post("/projects/{pid}/doors/rescan")
def rescan_doors(pid: int, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    """Re-read door tags from every plan on the project."""
    p = _own_project(pid, db, cu)
    total, plans = 0, 0
    for d in p.drawings:
        if d.status in ("uploaded", "processing"):
            continue
        n = sync_doors_from_drawing(db, d)
        if n:
            plans += 1
        total += n
    return {"doors": total, "plans": plans}


@router.put("/projects/{pid}/doors/bulk")
def bulk_update_doors(pid: int, payload: BulkDoorsIn, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    """Give many doors a set or a type at once: a list of ids, or every door with no type."""
    _own_project(pid, db, cu)
    if payload.set_id is not None and not db.query(models.HardwareSet).get(payload.set_id):
        raise HTTPException(404, "Set not found")
    if payload.door_type_id is not None:
        _own_door_type(payload.door_type_id, db, cu)
    q = db.query(models.Door).filter(models.Door.project_id == pid)
    if payload.untyped:
        q = q.filter(models.Door.door_type_id.is_(None))
    elif payload.door_ids:
        q = q.filter(models.Door.id.in_(payload.door_ids))
    else:
        return {"updated": 0}
    n = 0
    for d in q.all():
        if payload.clear_set:
            d.set_id = None
        elif payload.set_id is not None:
            d.set_id = payload.set_id
        if payload.door_type_id is not None:
            d.door_type_id = payload.door_type_id
        n += 1
    db.commit()
    return {"updated": n}


# ── Door schedule import ──────────────────────────────────────────────────────
_COL_ALIASES = {
    "code": ("door type", "door ref", "door no", "type", "ref"),
    "description": ("location description", "description", "location", "door description"),
    "width": ("width", "structural ope width", "w"),
    "height": ("height", "structural ope height", "h"),
    "fire": ("fire rating", "fire", "fd rating"),
    "acoustic": ("db rating", "acoustic", "acoustic rating", "sound"),
    "ironmongery": ("ironmongery", "hardware", "ironmongery set"),
    "floor": ("floor", "level"),
}


def _match_col(header: str) -> Optional[str]:
    h = re.sub(r"\s+", " ", (header or "").strip().lower())
    if not h:
        return None
    for key, names in _COL_ALIASES.items():
        if h in names:
            return key
    for key, names in _COL_ALIASES.items():
        if any(h.startswith(n) for n in names if len(n) > 3):
            return key
    return None


def _int_or_none(v):
    try:
        s = re.sub(r"[^\d.]", "", str(v))
        return int(float(s)) if s else None
    except ValueError:
        return None


@router.post("/projects/{pid}/doors/import-schedule")
async def import_door_schedule(pid: int, file: UploadFile = File(...),
                               db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    """
    An architect's door schedule (.xlsx). One row per door type (DT-01,
    Room Entrance Door, 1010 x 2135, FD30s, 37dB) fills in the door types;
    one row per door (D01-001 …) also creates the doors.
    """
    _own_project(pid, db, cu)
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Upload the schedule as an Excel file (.xlsx)")
    import openpyxl
    data = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(400, "Could not read that Excel file")

    cache = _type_cache(db, pid)
    types_added = types_updated = doors_added = 0
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        header_i, cols = None, {}
        for i, row in enumerate(rows[:40]):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if not any(re.fullmatch(r"door\s*(type|ref|no\.?)", c.lower()) for c in cells):
                continue
            cols = {}
            for j, c in enumerate(cells):
                k = _match_col(c)
                if k and k not in cols:
                    cols[k] = j
            # a "Width / Height" sub-header on the next row
            if i + 1 < len(rows):
                for j, c in enumerate(rows[i + 1]):
                    k = _match_col(str(c) if c is not None else "")
                    if k in ("width", "height") and k not in cols:
                        cols[k] = j
            header_i = i
            break
        if header_i is None or "code" not in cols:
            continue
        seen_types: set = set()
        for row in rows[header_i + 1:]:
            get = lambda k: (str(row[cols[k]]).strip() if k in cols and cols[k] < len(row) and row[cols[k]] is not None else "")
            raw = get("code")
            if not raw:
                continue
            code = _norm_code(raw)
            if not re.fullmatch(r"[A-Z]{1,4}-?\d{1,4}(?:-\d{1,2})?[A-Z]?", code) and not re.match(r"^D\d{2}-\d{3}$", code):
                continue
            desc = get("description")
            if desc.lower() in ("", "n/a", "-"):
                desc = ""
            per_door = bool(re.match(r"^D\d{2}-\d{3}$", code)) or bool(re.match(r"^D_[A-Z0-9.]+$", code))
            type_code = code[:3] if per_door and code.startswith("D") and "-" in code else code
            # "DT-05-01" / "DT-05-02" are size variants of DT-05; when the plans
            # only tag DT-05, fold them into that one type and keep the sizes.
            vm = re.fullmatch(r"(DT-\d{2,3})-(\d{2})", type_code)
            variant = None
            if vm:
                variant = type_code
                type_code = vm.group(1)
            dt = cache.get(type_code)
            if dt is None:
                dt = models.DoorType(project_id=pid, code=type_code, status="decide", sort_order=_code_sort(type_code)[1])
                db.add(dt); db.flush(); cache[type_code] = dt; types_added += 1
            elif type_code not in seen_types:
                types_updated += 1
            seen_types.add(type_code)
            if variant:
                base_desc = re.sub(r"\s*\(.*?\)?\s*$", "", desc).strip() if "(" in desc else desc
                if base_desc and not dt.description: dt.description = base_desc
                fire = get("fire")
                if fire and fire.lower() not in ("n/a", "-") and not dt.fire_rating: dt.fire_rating = fire
                ac = get("acoustic")
                if ac and ac.lower() not in ("n/a", "-", "none") and not dt.acoustic: dt.acoustic = ac
                w, h = _int_or_none(get("width")), _int_or_none(get("height"))
                if w and not dt.width: dt.width = w
                if h and not dt.height: dt.height = h
                line = f"{variant}: {desc}" + (f" {w} x {h}" if w and h else "")
                if line not in (dt.spec_text or ""):
                    dt.spec_text = ((dt.spec_text or "").rstrip() + "\n" + line).strip()
                continue
            if not per_door or not dt.description:
                if desc: dt.description = desc
                fire = get("fire")
                if fire and fire.lower() not in ("n/a", "-"): dt.fire_rating = fire
                ac = get("acoustic")
                if ac and ac.lower() not in ("n/a", "-", "none"): dt.acoustic = ac
                w, h = _int_or_none(get("width")), _int_or_none(get("height"))
                if w: dt.width = w
                if h: dt.height = h
                spec = get("ironmongery")
                if spec: dt.spec_text = spec
            if per_door:
                exists = db.query(models.Door).filter(models.Door.project_id == pid, models.Door.ref == code).first()
                if not exists:
                    db.add(models.Door(project_id=pid, door_type_id=dt.id, ref=code, floor=get("floor"),
                                       source="schedule", note=desc if per_door else ""))
                    doors_added += 1
    db.commit()
    return {"types_added": types_added, "types_updated": types_updated, "doors_added": doors_added}


# ── Individual doors ──────────────────────────────────────────────────────────
def _own_door(did: int, db: Session, cu) -> models.Door:
    d = db.query(models.Door).get(did)
    if not d:
        raise HTTPException(404, "Door not found")
    _own_project(d.project_id, db, cu)
    return d


def _door_single_out(db: Session, d: models.Door) -> DoorOut:
    dwg = db.query(models.Drawing).get(d.drawing_id) if d.drawing_id else None
    set_codes = {s.id: s.code for s in db.query(models.HardwareSet).all()}
    return _door_out(d, {dwg.id: (dwg.level or dwg.original_name)} if dwg else {}, set_codes)


@router.post("/projects/{pid}/doors", response_model=DoorOut, status_code=201)
def create_door(pid: int, payload: DoorIn, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    _own_project(pid, db, cu)
    if payload.door_type_id is not None:
        _own_door_type(payload.door_type_id, db, cu)
    ref = payload.ref.strip()
    if not ref:
        n = db.query(models.Door).filter(models.Door.project_id == pid, models.Door.source == "manual").count()
        ref = f"M-{n + 1:02d}"
    d = models.Door(project_id=pid, ref=ref, floor=payload.floor.strip(), handed=payload.handed,
                    door_type_id=payload.door_type_id, set_id=payload.set_id, note=payload.note, source="manual")
    db.add(d); db.commit(); db.refresh(d)
    return _door_single_out(db, d)


@router.put("/doors/{did}", response_model=DoorOut)
def update_door(did: int, payload: DoorIn, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    d = _own_door(did, db, cu)
    if payload.door_type_id is not None:
        _own_door_type(payload.door_type_id, db, cu)
    if payload.set_id is not None and not db.query(models.HardwareSet).get(payload.set_id):
        raise HTTPException(404, "Set not found")
    d.ref = payload.ref.strip() or d.ref; d.floor = payload.floor.strip(); d.handed = payload.handed
    d.door_type_id = payload.door_type_id; d.set_id = payload.set_id; d.note = payload.note
    db.commit(); db.refresh(d)
    return _door_single_out(db, d)


@router.delete("/doors/{did}", status_code=204)
def delete_door(did: int, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    d = _own_door(did, db, cu)
    db.delete(d); db.commit()


# ═════════════════════════════════════════════════════════════════════════════
# Schedule output
# ═════════════════════════════════════════════════════════════════════════════
from fastapi.responses import Response  # noqa: E402
from schedule_output import build_schedule, schedule_pdf, schedule_excel, picking_list_pdf  # noqa: E402


def _safe_name(project: models.Project, suffix: str) -> str:
    base = re.sub(r"[^\w\s\-]", "", project.name or "schedule").strip().replace(" ", "_") or "schedule"
    return f"{project.quote_no + '_' if project.quote_no else ''}{base}_{suffix}"


@router.get("/projects/{pid}/schedule")
def get_schedule(pid: int, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    """The schedule as data: sets with their doors and products, the product summary, and the checks."""
    p = _own_project(pid, db, cu)
    data = build_schedule(db, p, estimator=cu.name)
    for s in data["sets"]:
        for it in s["items"]:
            it["image_url"] = f"/api/files/products/{Path(it['image_path']).name}" if it["image_path"] else ""
            it.pop("image_path", None)
    return data


@router.get("/projects/{pid}/schedule/pdf")
def schedule_pdf_download(pid: int, priced: bool = False, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    p = _own_project(pid, db, cu)
    data = build_schedule(db, p, estimator=cu.name)
    if priced and not data["priced_ok"]:
        raise HTTPException(409, "Not every product on the schedule has a sell price")
    pdf = schedule_pdf(data, priced=priced)
    return Response(pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{_safe_name(p, "Schedule" + ("_Priced" if priced else ""))}.pdf"'})


@router.get("/projects/{pid}/schedule/excel")
def schedule_excel_download(pid: int, priced: bool = False, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    p = _own_project(pid, db, cu)
    data = build_schedule(db, p, estimator=cu.name)
    xlsx = schedule_excel(data, priced=priced and data["priced_ok"])
    return Response(xlsx, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{_safe_name(p, "Schedule")}.xlsx"'})


@router.get("/projects/{pid}/schedule/picking")
def picking_list_download(pid: int, db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    p = _own_project(pid, db, cu)
    data = build_schedule(db, p, estimator=cu.name)
    pdf = picking_list_pdf(data)
    return Response(pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{_safe_name(p, "Picking_List")}.pdf"'})


# ── Intec import: old schedules become the product and set library ───────────
@router.post("/sets/import-intec")
async def import_intec_schedule(file: UploadFile = File(...), create_project: bool = True,
                                db: Session = Depends(get_db), cu=Depends(auth.get_current_user)):
    """
    Upload an Intec schedule PDF (priced or unpriced). Its products and
    hardware sets are added to the library, and the job is kept as a
    door-schedule project so its sets are suggested on future jobs.
    """
    from intec_import import parse_intec_pdf, import_intec
    if not (file.filename or "").lower().endswith(".pdf"):
        raise HTTPException(400, "Upload the Intec schedule as a PDF")
    data = await file.read()
    try:
        parsed = parse_intec_pdf(data)
    except Exception as e:
        raise HTTPException(400, f"Could not read that as an Intec schedule: {e}")
    return import_intec(db, parsed, cu.id, create_project=create_project,
                        source_name=Path(file.filename or "").stem)
