"""
The ironmongery schedule itself: the numbers behind it (which doors get
which set, what that adds up to) and the documents that go out — the PDF
in the layout the office already sends, an Excel workbook, and a picking
list for the stores.
"""
import io
import re
from datetime import date
from pathlib import Path

from fpdf import FPDF
from fpdf.enums import XPos, YPos

import models

COMPANY = {
    "name": "MF Services",
    "address": ["26 Doughcloyne Court Industrial Estate", "Sarsfield Rd, Wilton Co Cork"],
    "phone": "PH 021 4348996",
    "web": "www.mfservices.ie",
}
ASSETS_DIR = Path(__file__).parent / "assets"
LOGO_PATH = ASSETS_DIR / "logo.png"          # office copy, if one has been dropped in
WEB_LOGO = ASSETS_DIR / "mf-logo.jpeg"        # the logo the web app shows; ships with the code
NAVY = (0, 56, 123); INK = (16, 25, 34); MUTED = (87, 100, 111); RULE = (196, 204, 212); SOFT = (242, 245, 247); BOX = (248, 250, 252)
PRODUCT_IMG_DIR = Path("uploads") / "products"

_FLOOR_ORDER = ["Basement", "Lower Ground", "Ground", "Mezzanine", "First", "Second", "Third", "Fourth",
                "Fifth", "Sixth", "Seventh", "Eighth", "Ninth", "Tenth", "Penthouse", "Roof"]


def _floor_rank(f):
    if f in _FLOOR_ORDER:
        return (0, _FLOOR_ORDER.index(f), f)
    m = re.match(r"Level (\d+)", f or "")
    return (1, int(m.group(1)), f) if m else (2, 0, f or "")


def _ref_key(ref: str):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", ref or "")]


# ── The numbers ───────────────────────────────────────────────────────────────
def build_schedule(db, project: models.Project, estimator: str = "") -> dict:
    """
    Everything the documents need, as plain data. Prices on the documents are
    the job's sell prices from the Cost Summary (product sell price, or the
    estimator's own figure on this job, less any discount). A product with no
    price at all leaves 'priced_ok' false.
    """
    types = {t.id: t for t in db.query(models.DoorType).filter(models.DoorType.project_id == project.id).all()}
    doors = db.query(models.Door).filter(models.Door.project_id == project.id).all()
    sets_by_id = {s.id: s for s in db.query(models.HardwareSet).all()}
    # Sell prices as set on the job's Cost Summary (actual S.P. after discounts); blank = product file.
    job_prices = {jp.product_id: jp for jp in db.query(models.JobPrice).filter_by(project_id=project.id).all()}

    def sell_price(p):
        jp = job_prices.get(p.id)
        cost = jp.cost if (jp and jp.cost is not None) else (p.cost or 0.0)
        sell = jp.sell if (jp and jp.sell is not None) else (p.sell if p.sell else cost)
        if jp:
            sell = sell * (1 - (jp.disc_a or 0) / 100) * (1 - (jp.disc_b or 0) / 100)
        return round(sell, 2) if sell else None

    groups: dict[int, list] = {}
    no_set, excluded = [], 0
    for d in doors:
        t = types.get(d.door_type_id)
        if d.set_id:
            eff = d.set_id
        elif t and t.status == "excluded":
            excluded += 1
            continue
        else:
            eff = t.set_id if t else None
        if not eff or eff not in sets_by_id:
            no_set.append(d)
            continue
        groups.setdefault(eff, []).append(d)

    sets_out, summary, priced_ok, photo_missing = [], {}, True, set()
    for sid, ds in sorted(groups.items(), key=lambda kv: _code_key(sets_by_id[kv[0]].code)):
        s = sets_by_id[sid]
        refs = sorted(ds, key=lambda d: (_floor_rank(d.floor or ""), _ref_key(d.ref)))
        items, per_door = [], 0.0
        for it in s.items:
            p = it.product
            price = sell_price(p)
            if price is None:
                priced_ok = False
            else:
                per_door += price * it.qty
            if not p.image_path or not Path(p.image_path).exists():
                photo_missing.add(p.sku)
            items.append({"product_id": p.id, "sku": p.sku, "name": p.name, "qty": it.qty, "unit": p.unit or "EACH",
                          "price": price, "value": (price * it.qty) if price is not None else None,
                          "image_path": p.image_path if p.image_path and Path(p.image_path).exists() else "",
                          "category": p.category or "Other"})
            agg = summary.setdefault(p.sku, {"sku": p.sku, "name": p.name, "unit": p.unit or "EACH",
                                             "category": p.category or "Other", "qty": 0, "price": price,
                                             "image_path": p.image_path if p.image_path and Path(p.image_path).exists() else ""})
            agg["qty"] += it.qty * len(refs)
        sets_out.append({
            "id": s.id, "code": s.code, "name": s.name, "fire_rated": bool(s.fire_rated),
            "items": items, "doors": len(refs),
            "door_refs": [{"ref": d.ref, "floor": d.floor or "", "handed": bool(d.handed),
                           "type_code": types[d.door_type_id].code if d.door_type_id in types else ""} for d in refs],
            "per_door": round(per_door, 2), "value": round(per_door * len(refs), 2),
        })

    summary_rows = sorted(summary.values(), key=lambda r: r["sku"].upper())
    for r in summary_rows:
        r["value"] = round(r["price"] * r["qty"], 2) if r["price"] is not None else None
    total = round(sum(s["value"] for s in sets_out), 2) if priced_ok else None

    types_to_decide = [t for t in types.values() if t.status == "decide"]
    checks = []
    if types_to_decide:
        checks.append({"level": "warn", "text": f"{len(types_to_decide)} door type{'s' if len(types_to_decide) != 1 else ''} still to decide: "
                                                + ", ".join(sorted(t.code for t in types_to_decide))})
    if no_set:
        checks.append({"level": "warn", "text": f"{len(no_set)} door{'s' if len(no_set) != 1 else ''} without a set, left off the schedule"})
    empty_sets = [s["code"] for s in sets_out if not s["items"]]
    if empty_sets:
        checks.append({"level": "warn", "text": "Sets with no products: " + ", ".join(empty_sets)})
    if not project.quote_no:
        checks.append({"level": "info", "text": "No quote number on the project"})
    if not project.client:
        checks.append({"level": "info", "text": "No client on the project"})
    if photo_missing:
        checks.append({"level": "info", "text": f"{len(photo_missing)} product{'s' if len(photo_missing) != 1 else ''} without a photo: "
                                                + ", ".join(sorted(photo_missing)[:8]) + (" …" if len(photo_missing) > 8 else "")})
    if not priced_ok and sets_out:
        checks.append({"level": "info", "text": "Some products have no cost yet (not landed in Cin7), so the priced version is off"})
    if sets_out and not checks:
        checks.append({"level": "ok", "text": "Every door has a set and every set has products"})

    return {
        "project": {"id": project.id, "name": project.name, "client": project.client or "",
                    "site": project.site or "", "quote_no": project.quote_no or "", "rep": project.rep or "",
                    "estimator": estimator, "date": date.today().strftime("%d/%m/%Y")},
        "sets": sets_out, "summary": summary_rows,
        "doors_scheduled": sum(s["doors"] for s in sets_out), "doors_no_set": len(no_set),
        "doors_excluded": excluded, "doors_total": len(doors),
        "priced_ok": priced_ok and bool(sets_out), "total": total, "checks": checks,
        "item_count": sum(r["qty"] for r in summary_rows),
    }


def _code_key(code: str):
    m = re.search(r"(\d+)", code or "")
    return (code.split(m.group(1))[0] if m else code, int(m.group(1)) if m else 0, code)


# ── PDF ───────────────────────────────────────────────────────────────────────
def _latin(s) -> str:
    """Core PDF fonts cover Latin-1; swap the odd dash or quote rather than crash."""
    s = "" if s is None else str(s)
    s = (s.replace("–", "-").replace("—", "-").replace("‘", "'").replace("’", "'")
          .replace("“", '"').replace("”", '"').replace("…", "...").replace(" ", " "))
    return s.encode("latin-1", "replace").decode("latin-1")


def _money(v) -> str:
    return "" if v is None else f"{v:,.2f}"


class SchedulePDF(FPDF):
    """House style: MF navy for bands and rules, the web logo top right, one photo per line."""

    def __init__(self, data: dict, priced: bool, title: str = ""):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.data = data
        self.priced = priced
        self.doc_title = title
        self.set_margins(18, 16, 18)
        self.set_auto_page_break(auto=True, margin=28)
        self.alias_nb_pages()
        self.set_title(f"{data['project']['name']} - {title or 'Ironmongery schedule'}")

    # Colour helpers -----------------------------------------------------
    def _ink(self):   self.set_text_color(*INK)
    def _muted(self): self.set_text_color(*MUTED)
    def _navy(self):  self.set_text_color(*NAVY)

    def _logo(self, x, y, w):
        path = LOGO_PATH if LOGO_PATH.exists() else (WEB_LOGO if WEB_LOGO.exists() else None)
        if path:
            self.image(str(path), x=x, y=y, w=w)
        else:
            self.set_font("Helvetica", "B", 20); self._navy(); self.set_xy(x, y); self.cell(w, 9, "MF Services", align="R"); self._ink()

    # Page furniture -----------------------------------------------------
    def header(self):
        p = self.data["project"]
        self._logo(x=148, y=12, w=44)
        self.set_font("Helvetica", "B", 10.5); self._ink()
        self.set_xy(18, 15); self.cell(120, 6, _latin(p["name"]))
        self.set_font("Helvetica", "", 8.5); self._muted()
        self.set_xy(18, 20.5); self.cell(120, 5, _latin(" | ".join(x for x in [self.doc_title or "Ironmongery schedule",
                                                                   f"Quote {p['quote_no']}" if p.get("quote_no") else ""] if x)))
        self.set_draw_color(*NAVY); self.set_line_width(0.6); self.line(18, 28, 192, 28); self.set_line_width(0.2)
        self._ink(); self.set_y(36)

    def footer(self):
        self.set_y(-24)
        self.set_draw_color(*RULE); self.line(18, self.get_y(), 192, self.get_y())
        self.set_font("Helvetica", "", 8); self._muted()
        y = self.get_y() + 1.5
        for i, line in enumerate([COMPANY["name"], *COMPANY["address"], f"{COMPANY['phone']} {COMPANY['web']}"]):
            self.set_xy(18, y + i * 3.6); self.cell(120, 3.6, line)
        self.set_xy(150, y); self.cell(42, 3.6, self.data["project"]["date"], align="R")
        self.set_xy(150, y + 3.6); self.cell(42, 3.6, f"Page {self.page_no()}/{{nb}}", align="R")
        self._ink()

    def band(self, left: str, mid: str = "", right: str = "", y: float = None):
        """The navy title band used for every section."""
        y = self.get_y() if y is None else y
        self.set_fill_color(*NAVY); self.rect(18, y, 174, 11, "F")
        self.set_text_color(255, 255, 255)
        self.set_font("Helvetica", "B", 11); self.set_xy(22, y + 2.5); self.cell(30, 6, _latin(left))
        if mid:
            self.set_font("Helvetica", "", 11); self.set_xy(48, y + 2.5); self.cell(100, 6, _latin(mid))
        if right:
            self.set_font("Helvetica", "", 9); self.set_xy(150, y + 2.5); self.cell(38, 6, _latin(right), align="R")
        self._ink(); self.set_y(y + 11)

    def _need(self, h: float, on_break=None):
        if self.get_y() + h > self.h - 28:
            self.add_page()
            if on_break: on_break()

    # Cover --------------------------------------------------------------
    def cover(self):
        self.add_page()
        p, d = self.data["project"], self.data
        self.band(self.doc_title or "Ironmongery Schedule", "", p["date"], y=40)
        self.set_y(58)
        left = [("Client", p["client"]), ("Project", p["name"]), ("Site", p["site"])]
        right = [("Quote no", p["quote_no"]), ("Your ref", d.get("your_ref", "")), ("Estimator", p["estimator"]), ("Rep", p["rep"])]
        for col, rows in ((18, left), (110, right)):
            y = 58
            for lbl, val in rows:
                self.set_font("Helvetica", "", 8.5); self._muted(); self.set_xy(col, y); self.cell(24, 6, lbl)
                self.set_font("Helvetica", "B", 10); self._ink(); self.set_xy(col + 24, y); self.cell(58, 6, _latin(val or ""))
                y += 8
        self.set_draw_color(*RULE); self.line(18, 92, 192, 92)
        self.set_font("Helvetica", "", 9.5); self._ink(); self.set_xy(18, 98)
        n_sets = len(d["sets"])
        self.cell(174, 6, f"{d['doors_scheduled']} doors across {n_sets} hardware set{'s' if n_sets != 1 else ''}, {d['item_count']} items.")
        y = 108
        if d.get("deliver_to"):
            self.set_font("Helvetica", "", 8.5); self._muted(); self.set_xy(18, y); self.cell(24, 5, "Deliver to")
            self.set_font("Helvetica", "", 10); self._ink(); self.set_xy(42, y); self.multi_cell(120, 5, _latin(d["deliver_to"]))
            y = self.get_y() + 4
        # contents
        self.set_font("Helvetica", "B", 8); self._muted(); self.set_xy(18, y + 4); self.cell(60, 5, "SETS ON THIS SCHEDULE")
        y += 10
        for s in d["sets"]:
            self.set_font("Helvetica", "B", 9); self._navy(); self.set_xy(18, y); self.cell(22, 5, _latin(s["code"]))
            self.set_font("Helvetica", "", 9); self._ink(); self.set_xy(40, y); self.cell(110, 5, _latin(s["name"]))
            self._muted(); self.set_xy(150, y); self.cell(42, 5, f"{s['doors']} door{'s' if s['doors'] != 1 else ''}", align="R")
            y += 5.5
            if y > self.h - 34: break
        self._ink()

    # Product rows -------------------------------------------------------
    def _cols(self):
        if self.priced:
            return [("Photo", 24, "L"), ("Code", 28, "L"), ("Description", 64, "L"), ("Qty", 10, "R"),
                    ("Price", 16, "R"), ("Unit", 14, "R"), ("Value", 18, "R")]
        return [("Photo", 26, "L"), ("Code", 30, "L"), ("Description", 88, "L"), ("Qty", 14, "R"), ("Unit", 16, "R")]

    def _table_header(self):
        self.set_fill_color(*SOFT); self.rect(18, self.get_y(), 174, 6, "F")
        self.set_font("Helvetica", "B", 7.5); self._muted(); self.set_xy(20, self.get_y() + 1)
        for name, w, align in self._cols():
            self.cell(w, 4, name, align=align)
        self._ink(); self.set_y(self.get_y() + 7)

    def _row(self, it: dict):
        cols = self._cols(); desc_w = cols[2][1]
        self.set_font("Helvetica", "", 8.5)
        lines = self.multi_cell(desc_w, 4, _latin(it["name"]), dry_run=True, output="LINES")
        h = max(20, 4 * len(lines) + 7)
        self._need(h, self._table_header)
        y = self.get_y(); x = 20
        # photo box
        self.set_fill_color(*BOX); self.set_draw_color(*RULE); self.rect(x, y + 1.5, 22, 16, "FD")
        img = it.get("image_path")
        if img:
            try: self.image(img, x=x + 1, y=y + 2.5, w=20, h=14, keep_aspect_ratio=True)
            except Exception: img = None
        if not img:
            self.set_font("Helvetica", "", 6.5); self._muted(); self.set_xy(x, y + 7.5); self.cell(22, 4, "No photo", align="C")
        x += cols[0][1]
        self.set_font("Helvetica", "B", 8.5); self._ink(); self.set_xy(x, y + 3); self.cell(cols[1][1], 4, _latin(it["sku"])); x += cols[1][1]
        self.set_font("Helvetica", "", 8.5); self.set_xy(x, y + 3); self.multi_cell(desc_w, 4, _latin(it["name"])); x += desc_w
        vals = [str(it["qty"])] + ([_money(it.get("price")), it["unit"], _money(it.get("value"))] if self.priced else [it["unit"]])
        for (n, w, align), v in zip(cols[3:], vals):
            self._muted() if n == "Unit" else self._ink()
            self.set_xy(x, y + 3); self.cell(w, 4, _latin(v), align=align); x += w
        self._ink(); self.set_draw_color(*RULE); self.line(18, y + h - 1, 192, y + h - 1)
        self.set_y(y + h)

    def _refs_panel(self, title: str, refs: list[str]):
        per_row, cw = 7, 24
        rows = max(1, (len(refs) + per_row - 1) // per_row)
        h = 9 + rows * 5 + 2
        self._need(h)
        top = self.get_y()
        self.set_fill_color(*SOFT); self.rect(18, top, 174, h, "F")
        self.set_font("Helvetica", "B", 8); self._navy(); self.set_xy(22, top + 2); self.cell(80, 5, title)
        self.set_font("Helvetica", "", 8.5); self._ink()
        x, y = 22, top + 8
        for i, r in enumerate(refs):
            if i and i % per_row == 0: x = 22; y += 5
            self.set_xy(x, y); self.cell(cw, 5, _latin(r)); x += cw
        if not refs:
            self._muted(); self.set_xy(22, y); self.cell(100, 5, "No doors on this set yet"); self._ink()
        self.set_y(top + h + 4)

    def set_page(self, s: dict):
        self.add_page()
        n = s["doors"]
        self.band(s["code"], s["name"], f"{n} door{'s' if n != 1 else ''}")
        self._table_header()
        for it in s["items"]:
            self._row(it)
        if not s["items"]:
            self.set_font("Helvetica", "I", 8.5); self._muted(); self.cell(0, 6, "No products in this set yet", new_y=YPos.NEXT); self._ink()
        if self.priced and s["items"]:
            self.set_font("Helvetica", "B", 8.5); y = self.get_y() + 1
            self.set_xy(110, y); self.cell(40, 5, f"{n} door{'s' if n != 1 else ''} @ {_money(s['per_door'])}", align="R")
            self.set_xy(150, y); self.cell(42, 5, _money(s["value"]), align="R")
            self.set_y(y + 7)
        else:
            self.set_y(self.get_y() + 3)
        self._refs_panel("Door references", [r["ref"] + ("h" if r["handed"] else "") for r in s["door_refs"]])

    # Summary / picking --------------------------------------------------
    def summary_page(self, picking: bool = False):
        self.add_page()
        rows = self.data["summary"]
        self.band("Picking list" if picking else "Product summary", "", f"{len(rows)} products | {self.data['item_count']} items")
        if picking:
            by_cat: dict[str, list] = {}
            for r in rows:
                by_cat.setdefault(r["category"], []).append(r)
            for cat in sorted(by_cat):
                self._need(14)
                self.set_font("Helvetica", "B", 8); self._navy(); self.set_y(self.get_y() + 2)
                self.cell(174, 5, _latin(cat), new_x=XPos.LMARGIN, new_y=YPos.NEXT); self._ink()
                self.set_draw_color(*RULE)
                for r in by_cat[cat]:
                    lines = self.multi_cell(104, 4, _latin(r["name"]), dry_run=True, output="LINES")
                    h = max(4 * len(lines), 4) + 2
                    self._need(h)
                    y = self.get_y()
                    self.rect(19, y + 0.5, 3.4, 3.4)
                    self.set_font("Helvetica", "", 8.5)
                    self.set_xy(26, y); self.cell(34, 4, _latin(r["sku"]))
                    self.set_xy(60, y); self.multi_cell(104, 4, _latin(r["name"]))
                    self.set_font("Helvetica", "B", 9); self.set_xy(164, y); self.cell(12, 4, str(r["qty"]), align="R")
                    self.set_font("Helvetica", "", 8.5); self._muted(); self.set_xy(176, y); self.cell(16, 4, _latin(r["unit"]), align="R"); self._ink()
                    self.set_y(y + h)
            return
        self._table_header()
        for r in rows:
            self._row(r)
        if self.priced:
            self.set_font("Helvetica", "B", 9.5); y = self.get_y() + 2
            self.set_xy(110, y); self.cell(40, 6, "Total", align="R")
            self._navy(); self.set_xy(150, y); self.cell(42, 6, _money(self.data["total"]), align="R"); self._ink()

    def notes_page(self):
        self.add_page()
        self.band("Notes")
        self.set_draw_color(*RULE)
        y = self.get_y() + 6
        while y < self.h - 40:
            self.line(18, y, 192, y); y += 8


def schedule_pdf(data: dict, priced: bool = False, summary: bool = True) -> bytes:
    pdf = SchedulePDF(data, priced=priced and data.get("priced_ok", False))
    pdf.cover()
    for s in data["sets"]:
        pdf.set_page(s)
    if summary:
        pdf.summary_page()
    pdf.notes_page()
    return bytes(pdf.output())


def picking_list_pdf(data: dict) -> bytes:
    packing = bool(data.get("deliver_to") or data.get("your_ref"))
    pdf = SchedulePDF(data, priced=False, title="Packing List" if packing else "Picking List")
    pdf.cover()
    if packing:
        pdf.add_page()
        pdf.band("Doors in this delivery", "", f"{data['doors_scheduled']} doors")
        pdf.set_y(pdf.get_y() + 3)
        for s in data["sets"]:
            n = s["doors"]
            pdf._refs_panel(f"{s['code']}  {s['name']}  ({n} door{'s' if n != 1 else ''})",
                            [r["ref"] + ("h" if r["handed"] else "") for r in s["door_refs"]])
    pdf.summary_page(picking=True)
    return bytes(pdf.output())


def schedule_excel(data: dict, priced: bool = False, summary: bool = True) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    priced = priced and data.get("priced_ok", False)
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E6E3DC")
    thin = Side(style="thin", color="999999")
    p = data["project"]

    def header(ws, cols):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font = bold; c.fill = head_fill; c.border = Border(bottom=thin)

    def widths(ws, ws_widths):
        for i, w in enumerate(ws_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # Sets
    ws = wb.active; ws.title = "Schedule"
    ws.append([f"{p['name']}"]); ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Quote {p['quote_no']}" if p["quote_no"] else "", p["client"], p["date"]])
    ws.append([])
    for s in data["sets"]:
        ws.append([f"Hardware Set Ref: {s['code']}", s["name"], f"{s['doors']} doors"])
        ws[ws.max_row][0].font = bold; ws[ws.max_row][1].font = bold
        header(ws, ["Product Code", "Description", "Qty per door", "Unit"] + (["Price", "Value per door"] if priced else []) + ["Total qty"])
        for it in s["items"]:
            ws.append([it["sku"], it["name"], it["qty"], it["unit"]] + ([it["price"], it["value"]] if priced else []) + [it["qty"] * s["doors"]])
        ws.append(["Door Reference", ", ".join(r["ref"] + ("h" if r["handed"] else "") for r in s["door_refs"])])
        ws[ws.max_row][0].font = bold
        ws[ws.max_row][1].alignment = Alignment(wrap_text=True, vertical="top")
        if priced:
            ws.append(["", "", "", "", f"{s['doors']} doors @", s["per_door"], s["value"]])
        ws.append([])
    widths(ws, [24, 70, 12, 8, 10, 14, 12])

    # Doors
    ws2 = wb.create_sheet("Doors")
    header(ws2, ["Door", "Floor", "Door type", "Handed", "Set", "Set name"])
    for s in data["sets"]:
        for r in s["door_refs"]:
            ws2.append([r["ref"], r["floor"], r["type_code"], "Yes" if r["handed"] else "", s["code"], s["name"]])
    widths(ws2, [14, 14, 12, 8, 10, 40])

    # Product summary (optional, like prices)
    if summary:
        ws3 = wb.create_sheet("Product summary")
        header(ws3, ["Product Code", "Description", "Category", "Qty", "Unit"] + (["Price", "Value"] if priced else []))
        for r in data["summary"]:
            ws3.append([r["sku"], r["name"], r["category"], r["qty"], r["unit"]] + ([r["price"], r["value"]] if priced else []))
        if priced:
            ws3.append([]); ws3.append(["", "", "", "", "", "Total", data["total"]]); ws3[ws3.max_row][6].font = bold
        widths(ws3, [24, 70, 18, 8, 8, 10, 12])

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.read()
